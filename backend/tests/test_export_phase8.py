from __future__ import annotations

import uuid
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from io import BytesIO
from typing import AsyncGenerator
from urllib.parse import unquote

import pytest
import pytest_asyncio
from fastapi import HTTPException
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

from app.database import get_db
from app.main import app
from app.middleware.rbac import create_access_token
from app.models.daily_report import DailyReport
from app.models.project import Project, ProjectHealthStatus, ProjectStatus, ProjectTrack
from app.models.risk_alert import RiskAlert
from app.models.sprint import Sprint, SprintStatus
from app.models.user import User, UserRole
from app.services import export as export_fonts
from app.services.export.project_summary_excel import build_project_summary_workbook
from app.services.export.reports_excel import build_reports_workbook
from app.services.export.scores_pdf import build_scores_pdf
from tests.conftest import TEST_DATABASE_URL


@pytest.fixture(scope="session", autouse=True)
def ensure_export_font_available() -> None:
    if not export_fonts.FONT_PATH.exists():
        pytest.skip("Phase 8 PDF 字体缺失,请先运行 python scripts/fetch_export_font.py")


@pytest_asyncio.fixture
async def db_session() -> AsyncGenerator[AsyncSession, None]:
    engine = create_async_engine(TEST_DATABASE_URL, echo=False)
    Session = async_sessionmaker(engine, expire_on_commit=False, class_=AsyncSession)
    async with Session() as session:
        try:
            yield session
        finally:
            await session.rollback()
    await engine.dispose()


@pytest_asyncio.fixture
async def client(db_session: AsyncSession) -> AsyncGenerator[AsyncClient, None]:
    async def _get_test_db():
        yield db_session

    app.dependency_overrides[get_db] = _get_test_db
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as ac:
        yield ac
    app.dependency_overrides.clear()


def _headers(user: User) -> dict[str, str]:
    return {"Authorization": f"Bearer {create_access_token(str(user.id), user.role.value)}"}


async def _make_user(db: AsyncSession, *, role: UserRole, department: str, name: str) -> User:
    user = User(
        id=uuid.uuid4(),
        wechat_userid=f"export_{uuid.uuid4().hex[:12]}",
        name=name,
        department=department,
        job_title="工程师",
        role=role,
        is_active=True,
    )
    db.add(user)
    await db.flush()
    return user


async def _seed_export_data(db: AsyncSession) -> dict[str, object]:
    salt = uuid.uuid4().hex[:6]
    dev_dept = f"研发导出部-{salt}"
    buy_dept = f"采购导出部-{salt}"
    today = date.today()

    manager = await _make_user(db, role=UserRole.manager, department=dev_dept, name="导出经理")
    employee = await _make_user(db, role=UserRole.employee, department=dev_dept, name="研发员工")
    buyer = await _make_user(db, role=UserRole.employee, department=buy_dept, name="采购员工")
    admin = await _make_user(db, role=UserRole.admin, department="总经办", name="导出管理员")

    project = Project(
        id=uuid.uuid4(),
        name="Phase 8 导出测试项目",
        code=f"EX{salt.upper()}",
        track=ProjectTrack.software,
        status=ProjectStatus.active,
        current_stage=3,
        health_status=ProjectHealthStatus.yellow,
        health_score=82,
        budget_total=Decimal("120000"),
        budget_spent=Decimal("45000"),
        created_by=manager.id,
    )
    db.add(project)
    await db.flush()

    sprint = Sprint(
        id=uuid.uuid4(),
        project_id=project.id,
        sprint_number=1,
        goal="完成数据导出",
        start_date=today - timedelta(days=10),
        end_date=today + timedelta(days=4),
        planned_story_points=20,
        completed_story_points=12,
        health_score=80,
        status=SprintStatus.active,
        created_by=manager.id,
    )
    db.add(sprint)
    await db.flush()

    reports = [
        DailyReport(
            id=uuid.uuid4(),
            user_id=employee.id,
            report_date=today - timedelta(days=2),
            raw_input_text="日报导出联调完成",
            parsed_content={
                "tasks": "实现日报汇总",
                "progress": 80,
                "acceptance_criteria": "三张 Sheet 可打开",
                "support_needed": "无",
                "blocker": "无",
                "next_step": "补测试",
                "eta": str(today),
                "reviewer": "导出经理",
                "git_version": "abc1234",
            },
            pass_check=True,
            ai_score=92,
            ai_comment="质量良好",
            project_id=project.id,
            created_by=employee.id,
        ),
        DailyReport(
            id=uuid.uuid4(),
            user_id=employee.id,
            report_date=today - timedelta(days=1),
            raw_input_text="PDF 字体存在风险",
            parsed_content={
                "tasks": "实现评分 PDF",
                "progress": 60,
                "acceptance_criteria": "中文可抽取",
                "support_needed": "字体文件",
                "blocker": "字体缺失会导致乱码",
                "next_step": "运行 fetch_export_font.py",
                "eta": str(today + timedelta(days=1)),
                "reviewer": "导出经理",
                "git_version": "def5678",
            },
            pass_check=False,
            ai_score=76,
            ai_comment="需要关注字体部署",
            project_id=project.id,
            created_by=employee.id,
        ),
        DailyReport(
            id=uuid.uuid4(),
            user_id=buyer.id,
            report_date=today - timedelta(days=1),
            raw_input_text="采购导出测试",
            parsed_content={
                "tasks": "采购物料",
                "progress": 100,
                "acceptance_criteria": "物料到货",
                "support_needed": "无",
                "blocker": "",
                "next_step": "入库",
            },
            pass_check=True,
            ai_score=88,
            ai_comment="推进顺利",
            created_by=buyer.id,
        ),
    ]
    db.add_all(reports)
    await db.flush()

    risk = RiskAlert(
        id=uuid.uuid4(),
        report_id=reports[1].id,
        user_id=employee.id,
        alert_type="blocker",
        description="字体缺失会导致 PDF 中文乱码",
        status="unresolved",
        days_unresolved=3,
        created_by=manager.id,
    )
    db.add(risk)
    await db.flush()

    return {
        "admin": admin,
        "manager": manager,
        "employee": employee,
        "buyer": buyer,
        "project": project,
        "dev_dept": dev_dept,
        "buy_dept": buy_dept,
        "today": today,
    }


@pytest.mark.asyncio
async def test_build_reports_workbook_returns_three_sheets(db_session: AsyncSession):
    seed = await _seed_export_data(db_session)
    today = seed["today"]

    workbook = load_workbook(
        await build_reports_workbook(
            db_session,
            start_date=today - timedelta(days=7),
            end_date=today,
        )
    )

    assert workbook.sheetnames == ["汇总", "每日明细", "部门小结"]


@pytest.mark.asyncio
async def test_build_reports_workbook_filters_by_department(db_session: AsyncSession):
    seed = await _seed_export_data(db_session)
    today = seed["today"]

    workbook = load_workbook(
        await build_reports_workbook(
            db_session,
            start_date=today - timedelta(days=7),
            end_date=today,
            department=seed["dev_dept"],
        )
    )
    detail = workbook["每日明细"]
    departments = [detail.cell(row=row, column=3).value for row in range(2, detail.max_row + 1)]

    assert departments
    assert set(departments) == {seed["dev_dept"]}


@pytest.mark.asyncio
async def test_build_scores_pdf_renders_chinese_without_tofu(db_session: AsyncSession):
    seed = await _seed_export_data(db_session)
    month = seed["today"].strftime("%Y-%m")

    pdf = await build_scores_pdf(db_session, month=month, department=seed["dev_dept"])
    text = _extract_pdf_text(pdf)

    assert "评分报告" in text
    assert "平均分" in text
    assert "■" not in text


@pytest.mark.asyncio
async def test_build_project_summary_workbook_404_when_deleted(db_session: AsyncSession):
    seed = await _seed_export_data(db_session)
    project = seed["project"]
    project.deleted_at = datetime.now(timezone.utc)
    await db_session.flush()

    with pytest.raises(HTTPException) as exc:
        await build_project_summary_workbook(db_session, project_id=project.id)

    assert exc.value.status_code == 404
    assert "项目不存在或已删除" in exc.value.detail


@pytest.mark.asyncio
async def test_pdf_skipped_when_font_missing(db_session: AsyncSession, monkeypatch, tmp_path):
    monkeypatch.setattr(export_fonts, "FONT_PATH", tmp_path / "missing.ttf")

    with pytest.raises(RuntimeError) as exc:
        await build_scores_pdf(db_session, month=date.today().strftime("%Y-%m"))

    assert "fetch_export_font.py" in str(exc.value)


@pytest.mark.asyncio
async def test_export_phase8_endpoints_return_files(client: AsyncClient, db_session: AsyncSession):
    seed = await _seed_export_data(db_session)
    today = seed["today"]
    headers = _headers(seed["manager"])

    reports = await client.get(
        "/api/v1/export/reports",
        params={
            "format": "xlsx",
            "start_date": str(today - timedelta(days=7)),
            "end_date": str(today),
            "department": seed["dev_dept"],
        },
        headers=headers,
    )
    assert reports.status_code == 200
    assert reports.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "AI日报汇总" in unquote(reports.headers["content-disposition"])
    assert load_workbook(BytesIO(reports.content)).sheetnames == ["汇总", "每日明细", "部门小结"]

    scores = await client.get(
        "/api/v1/export/scores",
        params={"format": "pdf", "month": today.strftime("%Y-%m"), "department": seed["dev_dept"]},
        headers=headers,
    )
    assert scores.status_code == 200
    assert scores.headers["content-type"].startswith("application/pdf")
    assert "评分报告" in unquote(scores.headers["content-disposition"])
    assert "评分报告" in _extract_pdf_text(BytesIO(scores.content))

    project = seed["project"]
    project_summary = await client.get(
        "/api/v1/export/project-summary",
        params={"format": "xlsx", "project_id": str(project.id)},
        headers=headers,
    )
    assert project_summary.status_code == 200
    assert project_summary.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert f"项目摘要_{project.code}" in unquote(project_summary.headers["content-disposition"])
    assert load_workbook(BytesIO(project_summary.content)).sheetnames == [
        "项目概况",
        "里程碑与 Sprint",
        "关联日报",
        "风险与卡点",
    ]


@pytest.mark.asyncio
async def test_export_phase8_endpoints_enforce_permissions(client: AsyncClient, db_session: AsyncSession):
    seed = await _seed_export_data(db_session)
    response = await client.get("/api/v1/export/reports", headers=_headers(seed["employee"]))

    assert response.status_code == 403


@pytest.mark.asyncio
async def test_export_phase8_endpoint_parameter_validation(client: AsyncClient, db_session: AsyncSession):
    seed = await _seed_export_data(db_session)
    headers = _headers(seed["manager"])

    invalid_month = await client.get("/api/v1/export/scores?month=2026-13", headers=headers)
    assert invalid_month.status_code == 422

    invalid_format = await client.get("/api/v1/export/reports?format=docx", headers=headers)
    assert invalid_format.status_code == 400


def _extract_pdf_text(pdf: BytesIO) -> str:
    pdf.seek(0)
    return "\n".join(page.extract_text() or "" for page in PdfReader(pdf).pages)
