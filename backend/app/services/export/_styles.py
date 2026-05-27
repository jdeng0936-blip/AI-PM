from __future__ import annotations

from collections.abc import Iterable, Sequence
from typing import Any

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

header_font = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
header_fill = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")
pass_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
fail_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
section_fill = PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid")
header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
data_align = Alignment(vertical="center", wrap_text=True)
center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
thin_border = Border(
    left=Side(style="thin", color="D1D5DB"),
    right=Side(style="thin", color="D1D5DB"),
    top=Side(style="thin", color="D1D5DB"),
    bottom=Side(style="thin", color="D1D5DB"),
)


def write_header(ws: Worksheet, headers: Sequence[str], *, row: int = 1) -> None:
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border


def write_row(
    ws: Worksheet,
    row: int,
    values: Sequence[Any],
    *,
    fill: PatternFill | None = None,
    center: bool = False,
) -> None:
    for col, value in enumerate(values, 1):
        cell = ws.cell(row=row, column=col, value=value)
        cell.alignment = center_align if center else data_align
        cell.border = thin_border
        if fill is not None:
            cell.fill = fill


def set_column_widths(ws: Worksheet, widths: Iterable[float]) -> None:
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, idx).column_letter].width = width


def write_section_title(ws: Worksheet, row: int, title: str, *, span: int = 4) -> None:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    cell = ws.cell(row=row, column=1, value=title)
    cell.font = Font(name="微软雅黑", bold=True, color="1F2937", size=12)
    cell.fill = section_fill
    cell.alignment = center_align
    cell.border = thin_border
