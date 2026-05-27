from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass
from datetime import date
from io import BytesIO
from statistics import mean

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from sqlalchemy import and_, func, select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.sql.elements import ColumnElement

from app.models.daily_report import DailyReport
from app.models.user import User
from app.services.export._styles import (
    fail_fill,
    pass_fill,
    set_column_widths,
    write_header,
    write_row,
    write_section_title,
)


@dataclass(frozen=True, slots=True)
class ReportExportRow:
    report: DailyReport
    user_name: str
    department: str


DETAIL_HEADERS = [
    "日期",
    "姓名",
    "部门",
    "AI评分",
    "状态",
    "今日任务",
    "完成进度",
    "验收标准",
    "所需支持",
    "核心卡点",
    "解决方案",
    "预计解决",
    "验收人",
    "Git版本",
    "AI评语",
]

DETAIL_WIDTHS = [12, 10, 14, 8, 8, 36, 10, 22, 18, 24, 24, 12, 10, 14, 34]

DEPT_HEADERS = ["部门", "提交人数", "平均分", "合格率", "卡点数"]
DEPT_WIDTHS = [18, 12, 12, 12, 12]


async def build_reports_workbook(
    db: AsyncSession,
    *,
    start_date: date,
    end_date: date,
    department: str | None = None,
) -> BytesIO:
    rows = await _fetch_report_rows(db, start_date=start_date, end_date=end_date, department=department)
    active_user_count = await _count_active_users(db, department=department)
    dept_summary = _build_department_summary(rows)
    summary = _build_summary(rows, active_user_count=active_user_count, start_date=start_date, end_date=end_date)
    top_users = _build_top_users(rows, limit=5)

    wb = Workbook()
    summary_ws = wb.active
    summary_ws.title = "汇总"
    _write_summary_sheet(summary_ws, summary, top_users, dept_summary, start_date, end_date, department)

    detail_ws = wb.create_sheet("每日明细")
    _write_detail_sheet(detail_ws, rows)

    dept_ws = wb.create_sheet("部门小结")
    _write_department_sheet(dept_ws, dept_summary)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


async def _fetch_report_rows(
    db: AsyncSession,
    *,
    start_date: date,
    end_date: date,
    department: str | None,
) -> list[ReportExportRow]:
    conditions: list[ColumnElement[bool]] = [
        DailyReport.report_date >= start_date,
        DailyReport.report_date <= end_date,
        DailyReport.deleted_at.is_(None),
    ]
    if department:
        conditions.append(User.department == department)

    stmt = (
        select(DailyReport, User.name, User.department)
        .join(User, DailyReport.user_id == User.id)
        .where(and_(*conditions))
        .order_by(DailyReport.report_date.desc(), DailyReport.ai_score.desc())
    )
    result = await db.execute(stmt)
    return [
        ReportExportRow(report=report, user_name=user_name, department=user_department or "未分配")
        for report, user_name, user_department in result.all()
    ]


async def _count_active_users(db: AsyncSession, *, department: str | None) -> int:
    conditions: list[ColumnElement[bool]] = [User.is_active.is_(True)]
    if department:
        conditions.append(User.department == department)

    result = await db.execute(select(func.count(User.id)).where(and_(*conditions)))
    return int(result.scalar_one() or 0)


def _build_summary(
    rows: list[ReportExportRow],
    *,
    active_user_count: int,
    start_date: date,
    end_date: date,
) -> dict[str, str | int | float]:
    scores = [row.report.ai_score for row in rows if row.report.ai_score is not None]
    checked = [row.report.pass_check for row in rows if row.report.pass_check is not None]
    pass_count = sum(1 for value in checked if value is True)
    reject_count = sum(1 for value in checked if value is False)
    submitted_pairs = {(row.report.user_id, row.report.report_date) for row in rows}
    period_days = max((end_date - start_date).days + 1, 1)
    expected_submissions = active_user_count * period_days
    blocker_count = sum(1 for row in rows if _has_blocker(row.report.parsed_content))

    return {
        "active_user_count": active_user_count,
        "report_count": len(rows),
        "submit_rate": _percentage(len(submitted_pairs), expected_submissions),
        "avg_score": round(mean(scores), 1) if scores else 0.0,
        "pass_rate": _percentage(pass_count, len(checked)),
        "reject_rate": _percentage(reject_count, len(checked)),
        "blocker_count": blocker_count,
    }


def _build_top_users(rows: list[ReportExportRow], *, limit: int) -> list[tuple[str, str, float, int]]:
    grouped: dict[tuple[str, str, str], list[int]] = defaultdict(list)
    for row in rows:
        if row.report.ai_score is None:
            continue
        grouped[(str(row.report.user_id), row.user_name, row.department)].append(row.report.ai_score)

    ranked = [
        (user_name, department, round(mean(scores), 1), len(scores))
        for (_, user_name, department), scores in grouped.items()
        if scores
    ]
    ranked.sort(key=lambda item: item[2], reverse=True)
    return ranked[:limit]


def _build_department_summary(rows: list[ReportExportRow]) -> list[dict[str, str | int | float]]:
    grouped: dict[str, list[ReportExportRow]] = defaultdict(list)
    for row in rows:
        grouped[row.department].append(row)

    summary: list[dict[str, str | int | float]] = []
    for department, dept_rows in grouped.items():
        scores = [row.report.ai_score for row in dept_rows if row.report.ai_score is not None]
        checked = [row.report.pass_check for row in dept_rows if row.report.pass_check is not None]
        pass_count = sum(1 for value in checked if value is True)
        summary.append(
            {
                "department": department,
                "submitter_count": len({row.report.user_id for row in dept_rows}),
                "avg_score": round(mean(scores), 1) if scores else 0.0,
                "pass_rate": _percentage(pass_count, len(checked)),
                "blocker_count": sum(1 for row in dept_rows if _has_blocker(row.report.parsed_content)),
            }
        )

    summary.sort(key=lambda item: (float(item["avg_score"]), int(item["submitter_count"])), reverse=True)
    return summary


def _write_summary_sheet(
    ws,
    summary: dict[str, str | int | float],
    top_users: list[tuple[str, str, float, int]],
    dept_summary: list[dict[str, str | int | float]],
    start_date: date,
    end_date: date,
    department: str | None,
) -> None:
    set_column_widths(ws, [22, 18, 16, 14, 14])
    write_section_title(ws, 1, "日报多维汇总", span=5)
    write_header(ws, ["指标", "数值"], row=3)

    metric_rows = [
        ("日期范围", f"{start_date} 至 {end_date}"),
        ("部门", department or "全部门"),
        ("活跃员工数", summary["active_user_count"]),
        ("日报提交数", summary["report_count"]),
        ("期间提交率", f"{summary['submit_rate']}%"),
        ("平均分", summary["avg_score"]),
        ("合格率", f"{summary['pass_rate']}%"),
        ("退回率", f"{summary['reject_rate']}%"),
        ("卡点数", summary["blocker_count"]),
    ]
    for row_idx, metric_values in enumerate(metric_rows, 4):
        write_row(ws, row_idx, metric_values)

    top_start = 15
    write_section_title(ws, top_start, "Top 5 高分员工", span=5)
    write_header(ws, ["姓名", "部门", "平均分", "日报数"], row=top_start + 1)
    for offset, top_values in enumerate(top_users, top_start + 2):
        write_row(ws, offset, top_values)

    chart_table_start = top_start + max(len(top_users), 1) + 5
    write_section_title(ws, chart_table_start, "部门小结图表数据", span=5)
    write_header(ws, DEPT_HEADERS, row=chart_table_start + 1)
    for offset, item in enumerate(dept_summary, chart_table_start + 2):
        write_row(
            ws,
            offset,
            [
                item["department"],
                item["submitter_count"],
                item["avg_score"],
                f"{item['pass_rate']}%",
                item["blocker_count"],
            ],
        )

    if dept_summary:
        chart = BarChart()
        chart.type = "bar"
        chart.style = 10
        chart.title = "部门平均分"
        chart.y_axis.title = "部门"
        chart.x_axis.title = "平均分"
        data = Reference(
            ws, min_col=3, min_row=chart_table_start + 1, max_row=chart_table_start + 1 + len(dept_summary)
        )
        categories = Reference(
            ws, min_col=1, min_row=chart_table_start + 2, max_row=chart_table_start + 1 + len(dept_summary)
        )
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(categories)
        chart.height = 7
        chart.width = 14
        ws.add_chart(chart, f"G{chart_table_start}")

    ws.freeze_panes = "A4"


def _write_detail_sheet(ws, rows: list[ReportExportRow]) -> None:
    set_column_widths(ws, DETAIL_WIDTHS)
    write_header(ws, DETAIL_HEADERS)
    for row_idx, row in enumerate(rows, 2):
        fill = pass_fill if row.report.pass_check is True else fail_fill if row.report.pass_check is False else None
        write_row(ws, row_idx, _detail_values(row), fill=fill)
    ws.freeze_panes = "A2"


def _write_department_sheet(ws, dept_summary: list[dict[str, str | int | float]]) -> None:
    set_column_widths(ws, DEPT_WIDTHS)
    write_header(ws, DEPT_HEADERS)
    for row_idx, item in enumerate(dept_summary, 2):
        write_row(
            ws,
            row_idx,
            [
                item["department"],
                item["submitter_count"],
                item["avg_score"],
                f"{item['pass_rate']}%",
                item["blocker_count"],
            ],
            center=True,
        )
    ws.freeze_panes = "A2"


def _detail_values(row: ReportExportRow) -> list[str | int | None]:
    report = row.report
    parsed = report.parsed_content or {}
    progress = parsed.get("progress")
    status = "合格" if report.pass_check is True else "退回" if report.pass_check is False else "未检查"
    return [
        str(report.report_date),
        row.user_name,
        row.department,
        report.ai_score,
        status,
        parsed.get("tasks", ""),
        f"{progress}%" if progress is not None else "",
        parsed.get("acceptance_criteria", ""),
        parsed.get("support_needed", ""),
        parsed.get("blocker", ""),
        parsed.get("next_step", ""),
        parsed.get("eta", ""),
        parsed.get("reviewer", ""),
        parsed.get("git_version", ""),
        report.ai_comment or "",
    ]


def _has_blocker(parsed_content: dict | None) -> bool:
    blocker = str((parsed_content or {}).get("blocker") or "").strip()
    return bool(blocker and blocker not in {"无", "暂无", "none", "None", "N/A", "n/a", "-"})


def _percentage(numerator: int, denominator: int) -> float:
    if denominator <= 0:
        return 0.0
    return round(numerator / denominator * 100, 1)
