from __future__ import annotations

from dataclasses import dataclass
from datetime import date
from io import BytesIO
from typing import Any
from uuid import UUID

from fastapi import HTTPException
from openpyxl import Workbook
from sqlalchemy import and_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.daily_report import DailyReport
from app.models.project import Project
from app.models.project_stage import ProjectStage
from app.models.risk_alert import RiskAlert
from app.models.sprint import Sprint
from app.models.user import User
from app.services.export._styles import (
    fail_fill,
    pass_fill,
    set_column_widths,
    write_header,
    write_row,
    write_section_title,
)


@dataclass(frozen=True, slots=True)
class ProjectReportRow:
    report: DailyReport
    user_name: str
    department: str


@dataclass(frozen=True, slots=True)
class ProjectRiskRow:
    risk: RiskAlert
    report_date: date
    user_name: str
    department: str


REPORT_HEADERS = [
    "日期",
    "姓名",
    "部门",
    "AI评分",
    "状态",
    "今日任务",
    "完成进度",
    "验收标准",
    "核心卡点",
    "解决方案",
    "AI评语",
]

REPORT_WIDTHS = [12, 10, 14, 8, 8, 36, 10, 22, 26, 26, 34]


async def build_project_summary_workbook(db: AsyncSession, *, project_id: UUID) -> BytesIO:
    project = await _fetch_project(db, project_id)
    owner_name = await _fetch_owner_name(db, project)
    stages = await _fetch_stages(db, project.id)
    sprints = await _fetch_sprints(db, project.id)
    report_rows = await _fetch_report_rows(db, project.id)
    risk_rows = await _fetch_risk_rows(db, project.id)

    wb = Workbook()
    overview_ws = wb.active
    overview_ws.title = "项目概况"
    _write_overview_sheet(overview_ws, project, owner_name)

    milestone_ws = wb.create_sheet("里程碑与 Sprint")
    _write_milestone_sheet(milestone_ws, stages, sprints)

    reports_ws = wb.create_sheet("关联日报")
    _write_reports_sheet(reports_ws, report_rows)

    risk_ws = wb.create_sheet("风险与卡点")
    _write_risk_sheet(risk_ws, risk_rows, report_rows)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


async def _fetch_project(db: AsyncSession, project_id: UUID) -> Project:
    result = await db.execute(select(Project).where(Project.id == project_id, Project.deleted_at.is_(None)))
    project = result.scalar_one_or_none()
    if project is None:
        raise HTTPException(status_code=404, detail="项目不存在或已删除")
    return project


async def _fetch_owner_name(db: AsyncSession, project: Project) -> str:
    if project.created_by is None:
        return ""
    result = await db.execute(select(User.name).where(User.id == project.created_by))
    return result.scalar_one_or_none() or ""


async def _fetch_stages(db: AsyncSession, project_id: UUID) -> list[ProjectStage]:
    result = await db.execute(
        select(ProjectStage).where(ProjectStage.project_id == project_id).order_by(ProjectStage.stage_number.asc())
    )
    return list(result.scalars().all())


async def _fetch_sprints(db: AsyncSession, project_id: UUID) -> list[Sprint]:
    result = await db.execute(
        select(Sprint)
        .where(Sprint.project_id == project_id)
        .order_by(Sprint.sprint_number.asc(), Sprint.start_date.asc())
    )
    return list(result.scalars().all())


async def _fetch_report_rows(db: AsyncSession, project_id: UUID) -> list[ProjectReportRow]:
    result = await db.execute(
        select(DailyReport, User.name, User.department)
        .join(User, DailyReport.user_id == User.id)
        .where(
            and_(
                DailyReport.project_id == project_id,
                DailyReport.deleted_at.is_(None),
            )
        )
        .order_by(DailyReport.report_date.desc(), DailyReport.ai_score.desc())
    )
    return [
        ProjectReportRow(report=report, user_name=user_name, department=user_department or "未分配")
        for report, user_name, user_department in result.all()
    ]


async def _fetch_risk_rows(db: AsyncSession, project_id: UUID) -> list[ProjectRiskRow]:
    result = await db.execute(
        select(RiskAlert, DailyReport.report_date, User.name, User.department)
        .join(DailyReport, RiskAlert.report_id == DailyReport.id)
        .join(User, RiskAlert.user_id == User.id)
        .where(
            and_(
                DailyReport.project_id == project_id,
                DailyReport.deleted_at.is_(None),
                RiskAlert.deleted_at.is_(None),
            )
        )
        .order_by(RiskAlert.created_at.desc())
    )
    return [
        ProjectRiskRow(
            risk=risk,
            report_date=report_date,
            user_name=user_name,
            department=user_department or "未分配",
        )
        for risk, report_date, user_name, user_department in result.all()
    ]


def _write_overview_sheet(ws, project: Project, owner_name: str) -> None:
    set_column_widths(ws, [18, 44])
    write_section_title(ws, 1, "项目概况", span=2)
    rows: list[tuple[str, Any]] = [
        ("名称", project.name),
        ("编码", project.code),
        ("track", _enum_value(project.track)),
        ("当前阶段", project.current_stage),
        ("健康状态", _enum_value(project.health_status)),
        ("健康度", project.health_score),
        ("负责人", owner_name),
        ("状态", _enum_value(project.status)),
        ("创建时间", _format_value(project.created_at)),
        ("计划上市", _format_value(project.planned_launch_date)),
        ("实际上市", _format_value(project.actual_launch_date)),
        ("预算总额", float(project.budget_total) if project.budget_total is not None else ""),
        ("预算已用", float(project.budget_spent) if project.budget_spent is not None else ""),
    ]
    for row_idx, values in enumerate(rows, 3):
        write_row(ws, row_idx, values)


def _write_milestone_sheet(ws, stages: list[ProjectStage], sprints: list[Sprint]) -> None:
    set_column_widths(ws, [10, 18, 14, 14, 14, 14, 14, 10, 12, 12, 48])
    write_section_title(ws, 1, "阶段与里程碑", span=11)
    stage_headers = [
        "阶段",
        "名称",
        "track",
        "计划开始",
        "计划结束",
        "实际开始",
        "实际结束",
        "进度",
        "健康状态",
        "Gate",
        "里程碑",
    ]
    write_header(ws, stage_headers, row=3)
    row_idx = 4
    for stage in stages:
        write_row(
            ws,
            row_idx,
            [
                stage.stage_number,
                stage.stage_name,
                _enum_value(stage.track),
                _format_value(stage.planned_start),
                _format_value(stage.planned_end),
                _format_value(stage.actual_start),
                _format_value(stage.actual_end),
                f"{stage.progress_pct}%",
                _enum_value(stage.health_status),
                "已通过" if stage.gate_passed else "未通过",
                _format_milestones(stage.milestones),
            ],
        )
        row_idx += 1

    row_idx += 2
    write_section_title(ws, row_idx, "Sprint", span=7)
    row_idx += 1
    write_header(ws, ["Sprint", "起始日期", "结束日期", "状态", "完成率", "燃尽情况", "健康分", "目标"], row=row_idx)
    row_idx += 1
    for sprint in sprints:
        write_row(
            ws,
            row_idx,
            [
                sprint.sprint_number,
                _format_value(sprint.start_date),
                _format_value(sprint.end_date),
                _enum_value(sprint.status),
                _completion_rate(sprint),
                _burndown_text(sprint),
                sprint.health_score,
                sprint.goal or "",
            ],
        )
        row_idx += 1
    ws.freeze_panes = "A4"


def _write_reports_sheet(ws, rows: list[ProjectReportRow]) -> None:
    set_column_widths(ws, REPORT_WIDTHS)
    write_header(ws, REPORT_HEADERS)
    for row_idx, row in enumerate(rows, 2):
        fill = pass_fill if row.report.pass_check is True else fail_fill if row.report.pass_check is False else None
        write_row(ws, row_idx, _report_values(row), fill=fill)
    ws.freeze_panes = "A2"


def _write_risk_sheet(ws, risk_rows: list[ProjectRiskRow], report_rows: list[ProjectReportRow]) -> None:
    set_column_widths(ws, [12, 12, 14, 16, 12, 14, 48, 28, 14])
    write_section_title(ws, 1, "风险预警历史", span=9)
    write_header(ws, ["日期", "姓名", "部门", "类型", "状态", "未解决天数", "描述", "物料/PO", "更新时间"], row=3)
    row_idx = 4
    for item in risk_rows:
        write_row(
            ws,
            row_idx,
            [
                _format_value(item.report_date),
                item.user_name,
                item.department,
                item.risk.alert_type,
                item.risk.status,
                item.risk.days_unresolved,
                item.risk.description,
                " / ".join(value for value in [item.risk.material_code, item.risk.po_number] if value),
                _format_value(item.risk.updated_at),
            ],
        )
        row_idx += 1

    row_idx += 2
    write_section_title(ws, row_idx, "日报卡点", span=7)
    row_idx += 1
    write_header(ws, ["日期", "姓名", "部门", "核心卡点", "解决方案", "预计解决", "AI评语"], row=row_idx)
    row_idx += 1
    for row in report_rows:
        parsed = row.report.parsed_content or {}
        blocker = str(parsed.get("blocker") or "").strip()
        if not _has_blocker(blocker):
            continue
        write_row(
            ws,
            row_idx,
            [
                _format_value(row.report.report_date),
                row.user_name,
                row.department,
                blocker,
                parsed.get("next_step", ""),
                parsed.get("eta", ""),
                row.report.ai_comment or "",
            ],
        )
        row_idx += 1
    ws.freeze_panes = "A4"


def _report_values(row: ProjectReportRow) -> list[str | int | None]:
    parsed = row.report.parsed_content or {}
    progress = parsed.get("progress")
    status = "合格" if row.report.pass_check is True else "退回" if row.report.pass_check is False else "未检查"
    return [
        str(row.report.report_date),
        row.user_name,
        row.department,
        row.report.ai_score,
        status,
        parsed.get("tasks", ""),
        f"{progress}%" if progress is not None else "",
        parsed.get("acceptance_criteria", ""),
        parsed.get("blocker", ""),
        parsed.get("next_step", ""),
        row.report.ai_comment or "",
    ]


def _format_milestones(milestones: list | None) -> str:
    if not milestones:
        return ""
    parts: list[str] = []
    for item in milestones:
        if not isinstance(item, dict):
            continue
        name = item.get("name") or "未命名"
        status = item.get("status") or "unknown"
        planned = item.get("planned_date") or "-"
        actual = item.get("actual_date") or "-"
        parts.append(f"{name}({status},计划:{planned},实际:{actual})")
    return "\n".join(parts)


def _completion_rate(sprint: Sprint) -> str:
    planned = sprint.planned_story_points or 0
    if planned <= 0:
        return "0%"
    return f"{round((sprint.completed_story_points or 0) / planned * 100, 1)}%"


def _burndown_text(sprint: Sprint) -> str:
    planned = sprint.planned_story_points or 0
    completed = sprint.completed_story_points or 0
    remaining = max(planned - completed, 0)
    return f"完成 {completed}/{planned} SP, 剩余 {remaining} SP"


def _has_blocker(blocker: str) -> bool:
    return bool(blocker and blocker not in {"无", "暂无", "none", "None", "N/A", "n/a", "-"})


def _enum_value(value: Any) -> str:
    return str(getattr(value, "value", value) or "")


def _format_value(value: Any) -> str:
    if value is None:
        return ""
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value)
